import unittest
from unittest.mock import patch

from docx import Document
from docx.shared import Pt

import cv_date_guard
import docx_adapter
import gemini_service
import pdf_generator
from models import CVAdaptation


class PdfFilenameTests(unittest.TestCase):
    def test_job_title_initials_are_in_pdf_name(self):
        self.assertEqual(
            pdf_generator.build_pdf_name("Zalando", "Data Analyst Recommerce"),
            "CV_Christian_Molina_Zalando_(DAR).pdf",
        )
        self.assertEqual(
            pdf_generator.build_pdf_name(
                "Workday",
                "Sr. Reporting & Analytics Consultant (FIN)",
            ),
            "CV_Christian_Molina_Workday_(SRACF).pdf",
        )

    def test_gender_markers_do_not_pollute_initials(self):
        self.assertEqual(
            pdf_generator.build_pdf_name(
                "Yepoda",
                "Performance Marketing Manager (m/f/d)",
            ),
            "CV_Christian_Molina_Yepoda_(PMM).pdf",
        )




class CvBulletStyleTests(unittest.TestCase):
    def test_banned_ai_style_word_is_rejected(self):
        problems = gemini_service._bullet_style_problems(
            "Orchestrated regional campaigns across multiple markets and channels."
        )
        self.assertTrue(any("orchestrated" in problem for problem in problems))

    def test_truncated_fragment_is_rejected(self):
        problems = gemini_service._bullet_style_problems(
            "Led cross functional analytics teams, translating business intent to."
        )
        self.assertTrue(any("colgante" in problem for problem in problems))

    def test_complete_natural_sentence_passes_style_guard(self):
        problems = gemini_service._bullet_style_problems(
            "Led regional analytics projects and built dashboards that improved decision making across EMEA markets."
        )
        self.assertEqual(problems, [])



class GeminiTargetedRepairTests(unittest.TestCase):
    def test_repairs_only_invalid_profile_and_title_with_gemini_text(self):
        bullets = [
            "Led regional analytics initiatives across EMEA markets and built reporting solutions that improved visibility, decision making and execution for commercial teams.",
            "Developed Power BI dashboards and performance reporting used across multiple markets, reducing manual work and improving access to consistent business insights.",
            "Managed cross functional data projects with commercial and technical stakeholders, aligning requirements and delivering practical solutions for regional business needs.",
            "Analyzed campaign and commercial performance using structured reporting and KPI frameworks, identifying opportunities and supporting data informed business decisions.",
        ]
        adaptation = CVAdaptation(
            nuevo_titulo="Senior Product Analytics Specialist",
            nuevo_perfil="X" * 778,
            nuevo_cargo_actual="Senior Product Data Analyst",
            nuevas_tareas=bullets,
            empresa_actual_sin_cambios="Stanley Black & Decker",
            fechas_actual_sin_cambios="(Aug 2024 – Jun 2026)",
        )

        gemini_profile = (
            "Senior product analytics professional with experience turning commercial and digital data into practical insights across EMEA markets. "
            "Skilled in Power BI, SQL, experimentation and performance analysis, with a strong record of partnering across business and technical teams. "
            "Experienced in building scalable reporting, improving visibility and supporting data informed decisions in international environments. "
            "Combines analytical depth with clear stakeholder communication and hands on execution."
        )
        gemini_title = "Senior Product Analytics Data Specialist"
        # Replacements must satisfy the same minimum as production (555 chars).
        gemini_profile += " Supports clear reporting and informed decisions for commercial teams."

        with patch.object(
            gemini_service,
            "_call_gemini_json",
            return_value={
                "nuevo_perfil": gemini_profile,
                "nuevo_titulo": gemini_title,
            },
        ):
            repaired = gemini_service._repair_invalid_cv_fields_with_gemini(
                adaptation,
                "Product analytics role focused on insights, experimentation and reporting.",
            )

        self.assertEqual(repaired.nuevo_perfil, gemini_profile)
        self.assertEqual(repaired.nuevo_titulo, gemini_title)
        self.assertEqual(repaired.nuevas_tareas, bullets)



class StableCvFlowTests(unittest.TestCase):
    def test_overlong_profile_and_short_title_use_one_targeted_repair(self):
        bullets = [
            "Led regional analytics initiatives across EMEA markets and built reporting solutions that improved visibility, decision making and execution for commercial and product teams.",
            "Developed Power BI dashboards and performance reporting used across multiple markets, reducing manual work and improving access to consistent business insights for regional teams.",
            "Managed cross functional data projects with commercial and technical stakeholders, aligning requirements and delivering practical solutions for regional business and analytics needs.",
            "Analyzed campaign and commercial performance using structured reporting and KPI frameworks, identifying opportunities and supporting data informed decisions across EMEA markets.",
        ]
        long_profile = "A" * 300 + ". " + "B" * 300 + ". " + "C" * 173 + "."

        full_generation = {
            "nuevo_titulo": "Senior Product Analytics Strategy Lead",
            "nuevo_perfil": long_profile,
            "nuevo_cargo_actual": "Senior Product Data Analyst",
            "nuevas_tareas": bullets,
            "empresa_actual_sin_cambios": "Stanley Black & Decker",
            "fechas_actual_sin_cambios": "(Aug 2024 – Jun 2026)",
        }
        targeted_title_repair = {
            "nuevo_titulo": "Senior Product Analytics Data Specialist",
        }

        with patch.object(
            gemini_service,
            "_call_gemini_json",
            side_effect=[full_generation, targeted_title_repair],
        ) as mocked:
            result = gemini_service.adapt_cv(
                "CV master text",
                "Product analytics role focused on reporting and experimentation.",
            )

        self.assertEqual(mocked.call_count, 2)
        self.assertEqual(result.nuevo_titulo, "Senior Product Analytics Data Specialist")
        self.assertEqual(len(result.nuevo_perfil), 603)
        self.assertTrue(result.nuevo_perfil.endswith("."))
        self.assertEqual(result.nuevas_tareas, bullets)

    def test_profile_fit_never_cuts_mid_sentence(self):
        long_profile = "A" * 300 + ". " + "B" * 300 + ". " + "C" * 173 + "."
        adaptation = CVAdaptation(
            nuevo_titulo="Senior Product Analytics Data Specialist",
            nuevo_perfil=long_profile,
            nuevo_cargo_actual="Senior Product Data Analyst",
            nuevas_tareas=[
                "Led regional analytics initiatives across EMEA markets and built reporting solutions that improved visibility, decision making and execution for commercial and product teams.",
                "Developed Power BI dashboards and performance reporting used across multiple markets, reducing manual work and improving access to consistent business insights for regional teams.",
                "Managed cross functional data projects with commercial and technical stakeholders, aligning requirements and delivering practical solutions for regional business and analytics needs.",
                "Analyzed campaign and commercial performance using structured reporting and KPI frameworks, identifying opportunities and supporting data informed decisions across EMEA markets.",
            ],
            empresa_actual_sin_cambios="Stanley Black & Decker",
            fechas_actual_sin_cambios="(Aug 2024 – Jun 2026)",
        )

        fitted = gemini_service._force_fix_adaptation(adaptation)

        self.assertEqual(len(fitted.nuevo_perfil), 603)
        self.assertTrue(fitted.nuevo_perfil.endswith("."))
        self.assertNotIn("C" * 10, fitted.nuevo_perfil)



class DocxExperienceAnchorTests(unittest.TestCase):
    def test_experienced_profile_does_not_match_experience_heading(self):
        doc = Document()
        doc.add_heading("Sobre mí", level=1)
        doc.add_paragraph(
            "Experienced Account Manager with a strong background in digital marketing."
        )
        doc.add_heading("Training & Certifications", level=1)
        training = doc.add_paragraph(
            "Stanley Black & Decker - SLP (Stanley Leadership Program) 2-year international leadership program"
        )
        training.style = "List Bullet"

        doc.add_heading("Experiencia", level=1)
        current = doc.add_heading(
            "Senior Product Manager, Stanley Black & Decker",
            level=2,
        )
        doc.add_paragraph("Fortune 500 American Manufacturer")
        doc.add_paragraph("(Aug 2024 – Jun 2026)")
        current_bullets = []
        for i in range(4):
            p = doc.add_paragraph(f"Original current bullet {i + 1}.")
            p.style = "List Bullet"
            current_bullets.append(p)

        doc.add_heading(
            "Digital Marketing Specialist EMEA ANZ, Stanley Black & Decker",
            level=2,
        )
        historical = []
        for i in range(3):
            p = doc.add_paragraph(f"Historical bullet {i + 1}.")
            p.style = "List Bullet"
            historical.append(p)

        role_idx = docx_adapter._find_current_role_heading_index(doc)

        self.assertEqual(doc.paragraphs[role_idx].text, current.text)
        self.assertNotEqual(doc.paragraphs[role_idx].text, training.text)

        new_tasks = [
            ("A" * 174) + ".",
            ("B" * 174) + ".",
            ("C" * 174) + ".",
            ("D" * 174) + ".",
        ]
        docx_adapter.update_current_role_bullets(doc, role_idx, new_tasks)

        self.assertEqual(training.text, "Stanley Black & Decker - SLP (Stanley Leadership Program) 2-year international leadership program")
        self.assertEqual([p.text for p in current_bullets], new_tasks)
        self.assertEqual(
            [p.text for p in historical],
            ["Historical bullet 1.", "Historical bullet 2.", "Historical bullet 3."],
        )

    def test_current_role_bullet_mismatch_aborts_without_deleting_history(self):
        doc = Document()
        doc.add_heading("Experiencia", level=1)
        doc.add_heading("Senior Product Manager, Stanley Black & Decker", level=2)
        doc.add_paragraph("Fortune 500 American Manufacturer")
        for i in range(3):
            p = doc.add_paragraph(f"Current bullet {i + 1}.")
            p.style = "List Bullet"
        doc.add_heading(
            "Digital Marketing Specialist EMEA ANZ, Stanley Black & Decker",
            level=2,
        )
        historical = doc.add_paragraph("Historical bullet must survive.")
        historical.style = "List Bullet"

        role_idx = docx_adapter._find_current_role_heading_index(doc)
        new_tasks = [
            ("A" * 174) + ".",
            ("B" * 174) + ".",
            ("C" * 174) + ".",
            ("D" * 174) + ".",
        ]

        with self.assertRaises(ValueError):
            docx_adapter.update_current_role_bullets(doc, role_idx, new_tasks)

        self.assertEqual(historical.text, "Historical bullet must survive.")


class CvTitleStyleTests(unittest.TestCase):
    def test_dangling_ampersand_title_is_rejected(self):
        problems = gemini_service._title_style_problems(
            "Account Manager, Digital Client Relations &"
        )
        self.assertTrue(problems)

    def test_internal_commas_are_not_restricted(self):
        for title in (
            "Senior Product Manager, AI Member Experiencia",
            "Marketing Operations Manager, Digital Growth",
            "Sales Account Manager, DACH, Strategic Focus",
            "Pricing, Revenue Operations, and Analytics",
        ):
            with self.subTest(title=title):
                self.assertEqual(
                    gemini_service._title_style_problems(title),
                    [],
                )



class ExperienceTaskFontSizeTests(unittest.TestCase):
    def test_all_experience_task_runs_use_same_ten_point_size(self):
        doc = Document()

        doc.add_heading("Training & Certifications", level=1)
        training = doc.add_paragraph()
        training.style = "List Bullet"
        training_run = training.add_run("Training bullet stays unchanged.")
        training_run.font.size = Pt(9)

        doc.add_heading("Experiencia", level=1)

        doc.add_heading(
            "Senior Performance Marketing Manager, Stanley Black & Decker",
            level=2,
        )
        recent = doc.add_paragraph()
        recent.style = "List Bullet"
        recent_run = recent.add_run("Recent role task.")
        recent_run.font.size = Pt(9)

        doc.add_heading(
            "Digital Marketing Specialist EMEA ANZ, Stanley Black & Decker",
            level=2,
        )
        mixed = doc.add_paragraph()
        mixed.style = "List Bullet"
        mixed_run_1 = mixed.add_run("Led Landing Page Optimization ")
        mixed_run_1.font.size = Pt(9)
        mixed_run_2 = mixed.add_run(
            "project: improved engagement and conversion across EMEA."
        )
        mixed_run_2.font.size = Pt(10)

        doc.add_heading(
            "Category Manager Amazon EMEA ANZ, Stanley Black & Decker",
            level=2,
        )
        historical = doc.add_paragraph()
        historical.style = "List Bullet"
        historical_run = historical.add_run("Historical task.")
        historical_run.font.size = Pt(10)

        doc.add_heading("Education", level=1)
        education_list = doc.add_paragraph()
        education_list.style = "List Bullet"
        education_run = education_list.add_run("Post Experience list stays unchanged.")
        education_run.font.size = Pt(9)

        docx_adapter._normalize_experience_task_font_sizes(doc)

        for paragraph in (recent, mixed, historical):
            self.assertTrue(paragraph.runs)
            self.assertEqual(
                [run.font.size.pt for run in paragraph.runs],
                [10.0] * len(paragraph.runs),
            )

        self.assertEqual(training_run.font.size.pt, 9.0)
        self.assertEqual(education_run.font.size.pt, 9.0)




class VisualLayoutGuardTests(unittest.TestCase):
    def test_teamblue_title_is_detected_as_three_lines(self):
        bad = "Marketing Operations Team Lead, AI & MarTech"
        good = "Growth Marketing Manager, Digital Performance"
        known_good = "Marketing Operations Manager, Digital Growth"

        self.assertEqual(gemini_service._estimated_title_lines(bad), 3)
        self.assertEqual(gemini_service._estimated_title_lines(good), 2)
        self.assertEqual(gemini_service._estimated_title_lines(known_good), 2)

    def test_shorter_teamblue_wording_fits_two_lines(self):
        fixed = "Marketing Operations Lead, AI & MarTech"
        self.assertEqual(len(fixed), 39)
        self.assertEqual(gemini_service._estimated_title_lines(fixed), 2)

    def test_synthflow_first_task_is_detected_as_three_lines(self):
        bad = (
            "Executed end-to-end regional marketing campaigns, adapting global "
            "strategies into effective Go-To-Market plans for EMEA and DACH. "
            "Focused on optimizing messaging across various growth channels."
        )
        good = (
            "Oversaw customer experience integration initiatives, enhanced "
            "landing pages and digital assets, and guaranteed precise tracking "
            "and dependable deployment across varied technology platforms."
        )

        self.assertEqual(gemini_service._estimated_bullet_lines(bad), 3)
        self.assertEqual(gemini_service._estimated_bullet_lines(good), 2)

    def test_validation_reports_visual_wrap_even_when_character_count_is_valid(self):
        bullets = [
            (
                "Executed end-to-end regional marketing campaigns, adapting global "
                "strategies into effective Go-To-Market plans for EMEA and DACH. "
                "Focused on optimizing messaging across various growth channels."
            ),
            (
                "Managed integrated digital, social and partnership campaigns, "
                "overseeing the full production cycle from creative briefing to "
                "delivery while improving conversion paths through structured testing."
            ),
            (
                "Developed robust performance tracking frameworks, monitoring "
                "campaign KPIs and customer response to identify optimization "
                "opportunities and improve attribution across regional growth channels."
            ),
            (
                "Partnered with sales and marketing teams to enhance lead quality "
                "and streamline funnel conversion processes, using data insights "
                "to inform practical growth priorities across EMEA markets."
            ),
        ]
        adaptation = CVAdaptation(
            nuevo_titulo="Marketing Operations Team Lead, AI & MarTech",
            nuevo_perfil=(
                "Growth marketing professional with experience across EMEA markets, "
                "combining campaign execution, analytics and digital optimization. "
                "Skilled in translating business priorities into measurable regional "
                "initiatives, building reporting frameworks and improving conversion "
                "paths through structured experimentation. Experienced in cross "
                "functional stakeholder management and scalable execution across "
                "complex international environments."
            ),
            nuevo_cargo_actual="Growth Marketing Manager",
            nuevas_tareas=bullets,
            empresa_actual_sin_cambios="Stanley Black & Decker",
            fechas_actual_sin_cambios="(Aug 2024 – Jun 2026)",
        )

        problems = gemini_service._validate_adaptation(adaptation)

        self.assertTrue(any("nuevo_titulo se estima en 3 líneas" in p for p in problems))
        self.assertTrue(any("bullet 1 se estima en 3 líneas" in p for p in problems))


if __name__ == "__main__":
    unittest.main()
