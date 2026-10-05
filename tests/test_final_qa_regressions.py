"""Reproduce failures found while validating the generated PDFs and live scraping."""
import tempfile
import os
import unittest
from pathlib import Path
from docx import Document
from models import ApplicationResult
from candidate_bible import load_candidate_bible
from application_agent.deterministic_answers import resolve_known_answer
from application_agent.account_access import _extract_verification_code
import scraper
import pdf_generator


class FinalQaRegressions(unittest.TestCase):
    def test_pdf_metadata_preserves_actual_canonical_filename(self):
        filename = pdf_generator.build_pdf_name('Example', 'Project Manager')
        result = ApplicationResult(ID='qa',company='Example',job_title='Project Manager',
            german_required='NO',cv_profile='',last_position='',responses=[],
            motivation_answer='',experience_answer='',pdf_name=filename)
        self.assertEqual(result.pdf_name,filename)
        self.assertIn('(PM)',result.model_dump()['pdf_name'])

    def test_pdf_filename_cannot_contain_path_separators(self):
        result = ApplicationResult(ID='qa',company='',job_title='',german_required='NO',
            cv_profile='',last_position='',responses=[],motivation_answer='',
            experience_answer='',pdf_name='bad/file\\name.pdf')
        self.assertNotIn('/',result.pdf_name)
        self.assertNotIn('\\',result.pdf_name)

    def test_local_cv_languages_and_latest_role_resolve_correctly(self):
        with tempfile.TemporaryDirectory() as folder:
            doc = Document()
            doc.add_heading('Sample Candidate | Project Manager',0)
            doc.add_heading('Experiencia',1)
            doc.add_heading('Project Manager, Example',2)
            doc.add_paragraph('(Ene 2020 – Dic 2021)')
            doc.add_heading('Habilidades',1)
            doc.add_paragraph('Power BI, Python')
            doc.add_heading('Idiomas',1)
            doc.add_paragraph('Español - Nativo')
            doc.add_paragraph('Inglés - C1 (Professional Working Proficiency)')
            path=Path(folder)/'cv.docx';doc.save(path)
            yaml=Path(folder)/'candidate.yaml';yaml.write_text('{}')
            bible=load_candidate_bible(yaml,cv_path=path)
            self.assertEqual(bible.get_path('languages.english.level'),'C1')
            self.assertEqual(bible.get_path('languages.english.notes'),'Professional Working Proficiency')
            for question in ('Are you fluent in English?','Are you fluent in Spanish?'):
                self.assertEqual(resolve_known_answer(question,{},bible).value,'Yes')
            self.assertEqual(resolve_known_answer('What is your current company?',{},bible).value,'Example')
            self.assertEqual(resolve_known_answer('What is your current job title?',{},bible).value,'Project Manager')

    def test_spanish_linkedin_metadata_is_normalized(self):
        for label,expected in [('Híbrido','Hybrid'),('Presencial','On-site'),('Remoto','Remote')]:
            with self.subTest(label=label):
                html=f'<script>{{"workplaceType":"{label}"}}</script>'
                self.assertEqual(scraper.extract_linkedin_work_format_from_html(html),expected)
                self.assertEqual(scraper.extract_work_format('Modalidad '+label),expected)

    def test_verification_code_requires_label_and_accepts_is_or_es(self):
        for text in ('Your verification code is 482731.','Código de verificación es 482731.'):
            self.assertEqual(_extract_verification_code(text),'482731')
        self.assertEqual(_extract_verification_code('Call 482731 for assistance'), '')
        self.assertEqual(_extract_verification_code('Code is 123456789.'), '')

class RenderedLayoutGuardTests(unittest.TestCase):
    def create_pdf(self, folder, profile_lines):
        import fitz
        from types import SimpleNamespace
        doc=fitz.open();page=doc.new_page()
        profile=[f"Perfil de experiencia línea {i}." for i in range(profile_lines)]
        tasks=[f"Gestiono la actividad {i} del proyecto." for i in range(4)]
        page.insert_text((70,70),"Sample Candidate | Gestión de proyectos\ny mejora de procesos",fontsize=12)
        for i,line in enumerate(profile):page.insert_text((70,130+i*15),line,fontsize=10)
        for i,task in enumerate(tasks):page.insert_text((70,300+i*25),task,fontsize=10)
        file=Path(folder)/'layout.pdf';doc.save(file);doc.close()
        a=SimpleNamespace(nuevo_titulo='Gestión de proyectos y mejora de procesos',
                          nuevo_perfil=' '.join(profile),nuevas_tareas=tasks)
        return file,a

    def test_real_pdf_with_six_profile_lines_passes(self):
        from pdf_layout_guard import validate_pdf_layout
        with tempfile.TemporaryDirectory() as folder:
            file,a=self.create_pdf(folder,6)
            self.assertEqual(validate_pdf_layout(file,a,'Sample Candidate'),[])

    def test_real_pdf_with_seven_profile_lines_triggers_repair(self):
        from pdf_layout_guard import validate_pdf_layout
        with tempfile.TemporaryDirectory() as folder:
            file,a=self.create_pdf(folder,7)
            problems=validate_pdf_layout(file,a,'Sample Candidate')
            self.assertEqual(problems,[{'field':'nuevo_perfil','reason':'rendered_line_count',
                                       'lines':7,'minimum':6,'maximum':6}])

    def test_missing_generated_text_fails_closed(self):
        from pdf_layout_guard import validate_pdf_layout
        with tempfile.TemporaryDirectory() as folder:
            file,a=self.create_pdf(folder,6)
            a.nuevo_perfil='Text that was not rendered in the PDF.'
            self.assertTrue(any(p['reason']=='generated_text_missing_or_duplicated'
                                for p in validate_pdf_layout(file,a,'Sample Candidate')))

    def test_duplicate_generated_profile_on_same_page_is_rejected(self):
        import fitz
        from pdf_layout_guard import validate_pdf_layout
        with tempfile.TemporaryDirectory() as folder:
            file,a=self.create_pdf(folder,6)
            with fitz.open(file) as doc:
                doc[0].insert_text((70,450), '\n'.join(
                    f'Perfil de experiencia línea {i}.' for i in range(6)),fontsize=10)
                doc.saveIncr()
            self.assertIn({'field':'nuevo_perfil','reason':'generated_text_missing_or_duplicated'},
                          validate_pdf_layout(file,a,'Sample Candidate'))

    def test_repeated_task_in_historical_role_does_not_reject_current_role(self):
        import fitz
        from pdf_layout_guard import validate_pdf_layout
        with tempfile.TemporaryDirectory() as folder:
            file,a=self.create_pdf(folder,6)
            with fitz.open(file) as doc:
                doc[0].insert_text((70,450),a.nuevas_tareas[0],fontsize=10)
                doc.saveIncr()
            self.assertEqual(validate_pdf_layout(file,a,'Sample Candidate'),[])

    def test_extra_page_is_rejected_even_if_adapted_fields_fit(self):
        import fitz
        from pdf_layout_guard import validate_pdf_layout
        with tempfile.TemporaryDirectory() as folder:
            file,a=self.create_pdf(folder,6)
            with fitz.open(file) as doc:
                doc.new_page().insert_text((70,70),'Historical experience overflow')
                doc.new_page().insert_text((70,70),'Habilidades')
                doc.saveIncr()
            self.assertIn({'field':'pagination','reason':'page_count','pages':3,'expected':2},
                          validate_pdf_layout(file,a,'Sample Candidate',expected_pages=2))

    def test_experience_on_second_page_is_rejected(self):
        import fitz
        from pdf_layout_guard import validate_pdf_layout
        with tempfile.TemporaryDirectory() as folder:
            file,a=self.create_pdf(folder,6)
            with fitz.open(file) as doc:
                doc.new_page().insert_text((70,70),'Historical experience overflow\nHabilidades')
                doc.saveIncr()
            self.assertIn({'field':'pagination','reason':'experience_overflows_first_page'},
                          validate_pdf_layout(file,a,'Sample Candidate',expected_pages=2))

    @unittest.skipUnless(os.getenv('RUN_PDF_RENDER_TESTS') == '1', 'Opt in to real LibreOffice rendering with local CV/fonts')
    def test_real_template_and_adaptation_keep_two_pages(self):
        import fitz
        import docx_adapter
        from types import SimpleNamespace
        from pdf_layout_guard import validate_pdf_layout
        from test_christian_cv import ChristianCvTests, SOURCE
        with tempfile.TemporaryDirectory() as folder:
            adapted=Path(folder)/'adapted.docx'
            a=ChristianCvTests().adaptation()
            docx_adapter.apply_cv_adaptation(str(SOURCE),str(adapted),a.model_dump())
            original_bytes=SOURCE.read_bytes()
            for path in (SOURCE,adapted):
                source=Document(path)
                pdf=pdf_generator.docx_to_pdf_via_libreoffice(str(path),Path(folder))
                fields=SimpleNamespace(nuevo_titulo=source.paragraphs[1].text.partition('|')[2].strip(),
                    nuevo_perfil=source.paragraphs[9].text,
                    nuevas_tareas=[source.paragraphs[i].text for i in (23,24,25,26)])
                self.assertEqual(validate_pdf_layout(pdf,fields,
                    source.paragraphs[1].text.partition('|')[0].strip(),expected_pages=2),[])
                with fitz.open(pdf) as rendered:
                    for paragraph in source.paragraphs[27:42]:
                        if paragraph.text.strip():
                            from pdf_layout_guard import _find_lines
                            self.assertEqual(_find_lines(rendered,paragraph.text)[0][0],0)
            self.assertEqual(SOURCE.read_bytes(),original_bytes)


if __name__=='__main__':unittest.main()
