"""Regressions against Christian's authoritative Word template, without Gemini/ADC."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
import cv_date_guard
import docx_adapter
import gemini_service
from candidate_bible import load_candidate_bible
from models import CVAdaptation

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'Christian_CV.docx'
PROFILE = ('Ingeniero Civil Industrial con experiencia en gestión de proyectos y mejora de procesos. '
           'Gestioné cronogramas, riesgos y recursos en Acme, coordinando clientes, proveedores y equipos internos para cumplir los plazos, la calidad y el alcance. '
           'Utilicé indicadores de desempeño para dar seguimiento a los proyectos y apoyar decisiones operacionales. '
           'Mi experiencia en Example Motors incluyó análisis de datos, dashboards y automatización de reportes con Excel y Power BI. '
           'Busco aportar esta base analítica a la coordinación de proyectos y la mejora continua de operaciones.')
TASKS = [
    'Gestioné proyectos multidisciplinarios en Acme, coordinando clientes, proveedores y equipos internos de logística, ventas, compras y servicio para cumplir plazos, calidad y alcance.',
    'Implementé mejoras en los procesos de ejecución de proyectos para optimizar la coordinación entre áreas, dar seguimiento a los compromisos y apoyar el cumplimiento del alcance definido.',
    'Gestioné cronogramas, riesgos y recursos de los proyectos, utilizando indicadores de desempeño para monitorear avances, detectar desviaciones y apoyar decisiones junto al equipo interno.',
    'Coordiné equipos multifuncionales y mantuve una comunicación efectiva con clientes, proveedores y áreas internas durante el ciclo de los proyectos para dar seguimiento a plazos y compromisos.',
]


class ChristianCvTests(unittest.TestCase):
    def adaptation(self):
        return CVAdaptation(nuevo_titulo='Gestión de proyectos y mejora de procesos',
            nuevo_perfil=PROFILE, nuevo_cargo_actual='Project Manager', nuevas_tareas=TASKS,
            empresa_actual_sin_cambios='Acme', fechas_actual_sin_cambios='(Ene 2020 – Dic 2021)')

    @unittest.skipUnless(SOURCE.exists(), "Place the local Christian_CV.docx to test the real template")
    def test_latest_role_from_real_cv(self):
        role = cv_date_guard.read_current_role(str(SOURCE))
        doc = Document(SOURCE)
        self.assertEqual(role.company, doc.paragraphs[20].text.partition(',')[2].strip())
        self.assertEqual(role.dates, doc.paragraphs[21].text.strip())

    @unittest.skipUnless(SOURCE.exists(), "Place the local Christian_CV.docx to test the real template")
    def test_real_word_adaptation_preserves_history_separators_and_photo(self):
        source = Document(SOURCE)
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder) / 'adapted.docx'
            a = self.adaptation()
            role = cv_date_guard.read_current_role(str(SOURCE))
            a.empresa_actual_sin_cambios = role.company
            a.fechas_actual_sin_cambios = role.dates
            self.assertEqual(gemini_service._validate_adaptation(a), [])
            docx_adapter.apply_cv_adaptation(str(SOURCE), str(out), a.model_dump())
            cv_date_guard.verify_immutable_dates(str(SOURCE), str(out))
            actual = Document(out)
            self.assertEqual(len(source.paragraphs), len(actual.paragraphs))
            changed = {1, 9, 20, 23, 24, 25, 26}
            for i, (before, after) in enumerate(zip(source.paragraphs, actual.paragraphs)):
                if i not in changed:
                    self.assertEqual(before._p.xml, after._p.xml, f'Paragraph {i} changed')
            self.assertEqual(actual.paragraphs[9].text, PROFILE)
            self.assertEqual([actual.paragraphs[i].text for i in [23,24,25,26]], TASKS)
            self.assertEqual(len(source.inline_shapes), len(actual.inline_shapes))
            self.assertEqual(source.sections[0]._sectPr.xml, actual.sections[0]._sectPr.xml)

    @unittest.skipUnless(SOURCE.exists(), "Place the local Christian_CV.docx to test the real template")
    def test_historical_date_change_is_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder) / 'bad.docx'
            doc = Document(SOURCE)
            doc.paragraphs[28].text = '(Ago 2024 – Jun 2026)'
            doc.save(out)
            with self.assertRaisesRegex(ValueError, 'fechas'):
                cv_date_guard.verify_immutable_dates(str(SOURCE), str(out))

    def test_repair_receives_master_cv_and_candidate_facts(self):
        a = self.adaptation()
        a.nuevo_perfil = 'Perfil corto.'
        truth = 'CV_MAESTRO: Acme y Example Motors. CANDIDATE_BIBLE_SOURCE_OF_TRUTH: Power BI y Python.'
        with patch.object(gemini_service, '_call_gemini_json', return_value={'nuevo_perfil': PROFILE}) as call:
            gemini_service._repair_invalid_cv_fields_with_gemini(a, 'Oferta con SQL', truth)
        self.assertIn(truth, call.call_args.args[1])
        self.assertIn('Spanish only', call.call_args.args[0])

    @unittest.skipUnless(SOURCE.exists(), "Place the local Christian_CV.docx to test candidate facts")
    def test_bible_contains_only_documented_candidate_facts(self):
        bible = load_candidate_bible(ROOT / 'candidate_bible.yaml', cv_path=SOURCE)
        self.assertEqual(bible.get_path('personal.full_name'), Document(SOURCE).paragraphs[1].text.split('|')[0].strip())
        self.assertEqual(bible.get_path('languages.english.level'), 'C1 (Professional Working Proficiency)')
        self.assertEqual(bible.get_path('personal.work_authorization'), '')
        self.assertNotIn('SQL', bible.get_path('skills.tools'))
        self.assertEqual(bible.get_path('professional.years_experience', ''), '')


if __name__ == '__main__':
    unittest.main()
