import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook
import excel_output as output


class ExcelOutputTests(unittest.TestCase):
    def test_url_reference_is_resolved_without_replacing_workbook_formula(self):
        output.ensure_workbook()
        book=load_workbook(self.path)
        sheet=book.active
        sheet['A3']='https://example.com/jobs/3'
        sheet['C3']='=A3'
        book.save(self.path);book.close()
        rows=output.read_rows(lookup_column='real_apply_url',lookup_value='https://example.com/jobs/3')
        self.assertEqual(rows[0]['real_apply_url'],'https://example.com/jobs/3')
        book=load_workbook(self.path)
        self.assertEqual(book.active['C3'].value,'=A3');book.close()

    def test_circular_url_reference_is_not_sent_as_url(self):
        output.ensure_workbook()
        book=load_workbook(self.path);sheet=book.active
        sheet['A2']='=C2';sheet['C2']='=A2'
        book.save(self.path);book.close()
        self.assertEqual(output.read_rows()[0]['real_apply_url'],'')
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "postulaciones.xlsx"
        self.env = patch.dict(os.environ, {"OUTPUT_EXCEL_PATH": str(self.path)})
        self.env.start()

    def tearDown(self):
        self.env.stop()
        self.temp.cleanup()

    def test_empty_workbook_created_and_can_be_read(self):
        self.assertEqual(output.read_rows(), [])
        self.assertTrue(self.path.exists())
        book = load_workbook(self.path)
        self.assertEqual(list(book.active.values)[0], tuple(output.HEADERS))
        book.close()

    def test_discovery_never_resets_reviewed_offer(self):
        output.write_row(output.RowInput(row={"Link": "https://example.com/1", "Status": "READY", "company": "Example"}))
        output.write_row(output.RowInput(row={"Link": "https://example.com/1", "Status": "NEW"}, preserve_existing=True))
        self.assertEqual(len(output.read_rows()), 1)
        self.assertEqual(output.read_rows()[0]["Status"], "READY")
        output.write_row(output.RowInput(row={"Link": "https://example.com/1", "Status": "DONE"}))
        self.assertEqual(output.read_rows()[0]["company"], "Example")
        self.assertEqual(output.read_rows()[0]["Status"], "DONE")

    def test_second_matching_key_and_literal_text(self):
        output.write_row(output.RowInput(row={"real_apply_url": "https://example.com/apply", "Status": "READY", "reason": "=1+1", "count": 0}, matching_columns=["real_apply_url"]))
        output.write_row(output.RowInput(row={"real_apply_url": "https://example.com/apply", "Status": "DONE"}, matching_columns=["real_apply_url"]))
        self.assertEqual(len(output.read_rows()), 1)
        self.assertEqual(output.read_rows()[0]["count"], 0)
        book = load_workbook(self.path)
        self.assertEqual(book.active["J2"].data_type, "s")
        book.close()

    def test_missing_match_key_does_not_add_row(self):
        with self.assertRaises(HTTPException):
            output.write_row(output.RowInput(row={"Status": "DONE"}))
        self.assertEqual(output.read_rows(), [])

    def test_live_reread_only_returns_requested_offer(self):
        for i in range(2):
            output.write_row(output.RowInput(row={"Link": str(i), "real_apply_url": f"https://example.com/{i}", "Status": "READY"}))
        rows = output.read_rows(lookup_column="real_apply_url", lookup_value="https://example.com/1")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Link"], "1")
        self.assertEqual(output.read_rows(lookup_column="real_apply_url", lookup_value="https://example.com/missing"), [])

    def test_structured_autofill_errors_are_stored_as_text(self):
        row = output.write_row(output.RowInput(row={"Link": "1", "reason": {"status": "manual_required", "detail": "Revisar"}}))
        self.assertEqual(json.loads(row["reason"])["status"], "manual_required")

    def test_workflows_are_disconnected_from_original_sheet(self):
        total = 0
        for path in Path(__file__).resolve().parents[1].glob("*.json"):
            if path.name.startswith("package"):
                continue
            data = json.loads(path.read_text(encoding="utf-8-sig"))
            for node in data.get("nodes", []):
                self.assertNotEqual(node.get("type"), "n8n-nodes-base.googleSheets")
                if node.get("parameters", {}).get("url", "").endswith("/excel-output/rows"):
                    total += 1
                    self.assertNotIn("credentials", node)
                    if "Snapshot" in node["name"]:
                        self.assertTrue(node["alwaysOutputData"])
        self.assertEqual(total, 20)

    def test_christian_flow_has_chile_search_and_live_lookup(self):
        path = Path(__file__).resolve().parents[1] / "n8n_christian_postulaciones.json"
        data = json.loads(path.read_text())
        self.assertFalse(data["active"])
        names = {node["name"] for node in data["nodes"]}
        self.assertEqual(len(names),len(data["nodes"]))
        nodes = {node['name']:node for node in data['nodes']}
        self.assertTrue(nodes['Call LinkedIn Search']['parameters']['url'].endswith('/linkedin-search-tasks'))
        self.assertEqual(nodes['Call LinkedIn Search']['parameters']['options']['timeout'],30000)
        self.assertEqual(nodes['Esperar búsqueda (10 s)']['parameters']['amount'],10)
        completion = data['connections']['¿Búsqueda terminada?']['main']
        self.assertEqual(completion[1][0]['node'],'Esperar búsqueda (10 s)')
        self.assertEqual({edge['node'] for edge in completion[0]},
                         {'Split ready_jobs (backend triage)','Avanzar siguiente search'})
        self.assertFalse(any("EXACT80" in name or "EXACT 80" in name for name in names))
        for name, ports in data["connections"].items():
            self.assertIn(name, names)
            for outputs in ports.values():
                for edges in outputs:
                    for edge in edges:
                        self.assertIn(edge["node"], names)
        search = next(node for node in data["nodes"] if node["name"] == "Search Params")
        url = next(field["value"] for field in search["parameters"]["assignments"]["assignments"] if field["name"] == "search_urls_raw")
        self.assertEqual(url, "https://www.linkedin.com/jobs/search/?location=Chile")
        lookup = next(node for node in data["nodes"] if node["name"] == "Releer estado en vivo")
        self.assertEqual(lookup["parameters"]["queryParameters"]["parameters"][0]["value"], "real_apply_url")
