"""Local output workbook for Christian's n8n workflows (never committed)."""
import fcntl
import os
import json
import re
from typing import Any
from contextlib import contextmanager
from pathlib import Path
from tempfile import NamedTemporaryFile

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

ROOT = Path(__file__).resolve().parent
HEADERS = ["Link", "Status", "real_apply_url", "company", "job_title", "added_at",
           "ID", "pdf_name", "download_url", "reason", "work_format",
           "autofill_status", "screenshot", "description", "Format", "Salary", "Revisar", "autofill"]
router = APIRouter(prefix="/excel-output", tags=["Excel local"])


def workbook_path():
    path = Path(os.environ.get("OUTPUT_EXCEL_PATH", "data/postulaciones.xlsx"))
    return path if path.is_absolute() else ROOT / path


@contextmanager
def _locked_book():
    path = workbook_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    # OS lock also coordinates multiple backend worker processes on Ubuntu.
    with path.with_suffix(".xlsx.lock").open("a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        book = None
        try:
            if path.exists():
                book = load_workbook(path)
                sheet = book["Postulaciones"]
            else:
                book = Workbook()
                sheet = book.active
                sheet.title = "Postulaciones"
                sheet.append(HEADERS)
                sheet.freeze_panes = "C2"
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="164E63")
                for i, header in enumerate(HEADERS, 1):
                    sheet.column_dimensions[get_column_letter(i)].width = (
                        48 if header in {"Link", "real_apply_url", "download_url", "reason"}
                        else 28
                    )
                _save(book, sheet, path)
            yield book, sheet, path
        finally:
            if book is not None:
                book.close()
            fcntl.flock(lock, fcntl.LOCK_UN)


def _save(book, sheet, path):
    sheet.auto_filter.ref = sheet.dimensions
    with NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as tmp:
        temporary = Path(tmp.name)
    try:
        book.save(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def ensure_workbook():
    with _locked_book():
        pass
    return workbook_path()


def _headers(sheet):
    return [cell.value for cell in sheet[1]]


def _url_cell_value(sheet, cell, visited=None):
    visited = set() if visited is None else visited
    if cell.coordinate in visited or len(visited) >= 10:
        return ''
    visited.add(cell.coordinate)
    value = cell.value
    if isinstance(value, str) and value.startswith('='):
        reference = re.fullmatch(r'=\s*\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)\s*', value)
        if reference:
            return _url_cell_value(sheet, sheet[reference[1] + reference[2]], visited)
        return cell.hyperlink.target if cell.hyperlink else ''
    return value if value is not None else ''


def _row(sheet, number, headers):
    row = {}
    for i, header in enumerate(headers, 1):
        if header:
            cell = sheet.cell(number, i)
            row[header] = (_url_cell_value(sheet, cell) if header in {'Link','real_apply_url'}
                           else cell.value if cell.value is not None else '')
    return row


class RowInput(BaseModel):
    row: dict[str, Any]
    matching_columns: list[str] = Field(default_factory=lambda: ["Link"])
    preserve_existing: bool = False


@router.get("/rows")
def read_rows(lookup_column: str | None = None, lookup_value: str | None = None):
    with _locked_book() as (_, sheet, __):
        headers = _headers(sheet)
        rows = [_row(sheet, n, headers) for n in range(2, sheet.max_row + 1)
                if any(sheet.cell(n, i).value is not None
                       for i in range(1, sheet.max_column + 1))]
        if lookup_column is not None:
            if lookup_column not in headers or lookup_value is None:
                raise HTTPException(422, "Filtro de fila inválido.")
            rows = [row for row in rows if str(row.get(lookup_column, "")) == lookup_value]
        return rows


@router.post("/rows")
def write_row(payload: RowInput):
    keys = payload.matching_columns
    if not keys or any(not str(payload.row.get(key) or "").strip() for key in keys):
        raise HTTPException(422, "La fila debe tener Link o real_apply_url para identificar la oferta.")
    with _locked_book() as (book, sheet, path):
        headers = _headers(sheet)
        for key in payload.row:
            if key not in headers:
                headers.append(key)
                sheet.cell(1, len(headers), key)
        number = None
        for n in range(2, sheet.max_row + 1):
            existing = _row(sheet, n, headers)
            if all(str(existing.get(key) or "").strip() == str(payload.row[key]).strip()
                   for key in keys):
                number = n
                if payload.preserve_existing:
                    return existing
                break
        number = number or sheet.max_row + 1
        for key, value in payload.row.items():
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            cell = sheet.cell(number, headers.index(key) + 1)
            cell.value = value
            # Scraped text must remain text, even if it starts with '='.
            if isinstance(value, str):
                cell.data_type = "s"
        _save(book, sheet, path)
        return _row(sheet, number, headers)


@router.get("/download")
def download_workbook():
    return FileResponse(ensure_workbook(), filename="postulaciones.xlsx",
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    print(ensure_workbook())
