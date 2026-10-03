"""Validate dates against the authoritative CV without rewriting its XML."""
from dataclasses import dataclass
import re
from docx import Document


@dataclass(frozen=True)
class CurrentRole:
    company: str
    dates: str


def _date_ranges(doc):
    return [p.text.strip() for p in doc.paragraphs
            if re.fullmatch(r'\([^()]*\d{4}\s*[–—-]\s*[^()]*\d{4}\)', p.text.strip())]


def read_current_role(path: str) -> CurrentRole:
    paragraphs = Document(path).paragraphs
    section = next((i for i, p in enumerate(paragraphs)
                    if p.text.strip().casefold() == 'experiencia'), None)
    if section is None:
        raise ValueError('No se encontró la sección Experiencia del CV maestro')
    for i in range(section + 1, len(paragraphs)):
        p = paragraphs[i]
        if p.style.name == 'Heading 1' and set(p.text.strip()) != {'_'}:
            break
        if p.style.name == 'Heading 2' and ',' in p.text:
            company = p.text.partition(',')[2].strip()
            following = next((q.text.strip() for q in paragraphs[i + 1:]
                              if q.text.strip()), '')
            if not re.fullmatch(r'\([^()]*\d{4}\s*[–—-]\s*[^()]*\d{4}\)', following):
                raise ValueError('El rol más reciente no contiene un rango de fechas válido')
            return CurrentRole(company, following)
    raise ValueError('No se encontró el rol más reciente del CV maestro')


def verify_immutable_dates(source_path: str, output_path: str) -> None:
    source = _date_ranges(Document(source_path))
    output = _date_ranges(Document(output_path))
    if not source or source != output:
        raise ValueError('GUARDARRAÍL: las fechas del CV adaptado no coinciden con el CV maestro')
